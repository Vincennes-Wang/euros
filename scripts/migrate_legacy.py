"""Import data from legacy/硬币收藏清单.xlsx (sheet 纪念币).

- Cached GOOGLETRANSLATE values (column 描述) -> description_zh in data/coins.json
- Collection status (column Collection) -> private/collection.json (gitignored)

Matching: same year and country, then best similarity of Feature vs title
(or legacy Description vs description_en). Each coin matches at most one row.

Usage: python scripts/migrate_legacy.py [path/to/xlsx]
"""
from __future__ import annotations

import json
import re
import sys
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from scraper import ecb  # noqa: E402
from scraper.scrape import COINS_JSON, write_if_changed  # noqa: E402

XLSX = ROOT / "legacy" / "硬币收藏清单.xlsx"
COLLECTION_JSON = ROOT / "private" / "collection.json"
SHEET = "纪念币"
MIN_SCORE = 0.6       # below this a pair is not a match
CONFIDENT_SCORE = 0.85


def norm(s: str | None) -> str:
    return " ".join(re.findall(r"[a-z0-9]+", (s or "").lower()))


def similarity(a: str | None, b: str | None) -> float:
    return SequenceMatcher(None, norm(a), norm(b)).ratio()


def read_rows(path: Path) -> list[dict]:
    ws = openpyxl.load_workbook(path, data_only=True)[SHEET]
    header = [c.value for c in ws[1]]
    col = {name: header.index(name) for name in ("Year", "Country", "Feature", "Collection", "Description", "描述")}
    rows = []
    for excel_row, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if cells[col["Year"]] is None:
            continue
        zh = cells[col["描述"]]
        rows.append({
            "row": excel_row,
            "year": int(cells[col["Year"]]),
            "country": str(cells[col["Country"]] or "").strip(),
            "feature": cells[col["Feature"]],
            "description": cells[col["Description"]],
            "zh": zh.strip() if isinstance(zh, str) and zh.strip() and not zh.startswith("#") else None,
            "owned": cells[col["Collection"]] is True,
        })
    return rows


def match(rows: list[dict], coins: list[dict]) -> tuple[dict[int, tuple[dict, float]], list[dict]]:
    """Return ({excel_row: (coin, score)}, unmatched rows)."""
    pairs = []
    for r in rows:
        country = ecb.normalize_country(r["country"])
        code = country[1] if country else None
        for c in coins:
            if c["year"] == r["year"] and c["country_code"] == code:
                score = max(similarity(r["feature"], c["title"]), similarity(r["description"], c["description_en"]))
                pairs.append((score, r["row"], c["id"], r, c))
    pairs.sort(key=lambda p: (-p[0], p[1], p[2]))
    matched: dict[int, tuple[dict, float]] = {}
    used: set[str] = set()
    for score, row, cid, _, coin in pairs:
        if score < MIN_SCORE or row in matched or cid in used:
            continue
        matched[row] = (coin, score)
        used.add(cid)
    return matched, [r for r in rows if r["row"] not in matched]


def main() -> int:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else XLSX
    rows = read_rows(path)
    coins = json.loads(COINS_JSON.read_text(encoding="utf-8"))
    matched, unmatched = match(rows, coins)
    by_row = {r["row"]: r for r in rows}

    for row, (coin, _) in matched.items():
        zh = by_row[row]["zh"]
        if zh:
            coin["description_zh"] = zh
    data = (json.dumps(coins, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    changed = write_if_changed(COINS_JSON, data)

    owned = sorted(coin["id"] for row, (coin, _) in matched.items() if by_row[row]["owned"])
    unmatched_owned = [
        {"row": r["row"], "year": r["year"], "country": r["country"], "feature": r["feature"]}
        for r in unmatched if r["owned"]
    ]
    collection = {
        "source": path.name,
        "exported_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
        "owned": owned,
        "unmatched_owned": unmatched_owned,
    }
    COLLECTION_JSON.parent.mkdir(parents=True, exist_ok=True)
    COLLECTION_JSON.write_text(json.dumps(collection, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    legacy_years = {r["year"] for r in rows}
    low = sorted((s, row) for row, (_, s) in matched.items() if s < CONFIDENT_SCORE)
    drifted = sorted(
        (similarity(by_row[row]["description"], coin["description_en"]), row)
        for row, (coin, _) in matched.items()
    )
    missing_zh = [c["id"] for c in coins if c["year"] in legacy_years and not c.get("description_zh")]

    print(f"legacy rows: {len(rows)}; matched: {len(matched)}; unmatched: {len(unmatched)}")
    print(f"description_zh written: {sum(1 for c in coins if c.get('description_zh'))} "
          f"(coins.json {'updated' if changed else 'unchanged'})")
    print(f"owned: {len(owned)} matched, {len(unmatched_owned)} unmatched -> {COLLECTION_JSON.relative_to(ROOT)}")
    print("\n## Unmatched legacy rows")
    for r in unmatched:
        print(f"  row {r['row']}: {r['year']} {r['country']} | {r['feature']}{'  [OWNED]' if r['owned'] else ''}")
    print(f"\n## Low-confidence matches (score < {CONFIDENT_SCORE})")
    for score, row in low:
        r, coin = by_row[row], matched[row][0]
        print(f"  row {row} ({score:.2f}): {r['feature']!r} -> {coin['id']} {coin['title']!r}")
    print("\n## English description differs from legacy (zh may translate older text; similarity < 0.9)")
    for score, row in drifted:
        if score < 0.9:
            print(f"  row {row} ({score:.2f}): {matched[row][0]['id']}")
    print(f"\n## Coins in {min(legacy_years)}-{max(legacy_years)} without description_zh: {len(missing_zh)}")
    for cid in missing_zh:
        print(f"  {cid}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
